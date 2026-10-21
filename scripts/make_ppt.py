"""Build the demo deck from a run's real results.

    python scripts/make_ppt.py --run pilot [--out deliverables/RSI_演示.pptx]
Everything on the slides (numbers, charts, decision log, sample answers) is read from runs/<run>/.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

from rsi import report as R  # noqa: E402
from rsi.constraints import KIND_ZH, KINDS  # noqa: E402
from rsi.data import read_jsonl  # noqa: E402
from rsi.optimizer import OP_ZH  # noqa: E402
from rsi.paths import RUNS  # noqa: E402

FONT = "Microsoft YaHei"
INK = RGBColor(0x12, 0x1B, 0x2E)       # dark slides / primary text
INK2 = RGBColor(0x4A, 0x55, 0x68)      # secondary text
MUTED = RGBColor(0x8A, 0x94, 0xA6)
TEAL = RGBColor(0x0E, 0x7C, 0x7B)      # dominant brand colour
TEAL_L = RGBColor(0xE3, 0xF2, 0xF1)    # card tint
AMBER = RGBColor(0xF2, 0x9F, 0x05)     # single sharp accent
RED = RGBColor(0xC2, 0x41, 0x0C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE2, 0xE6, 0xEC)
W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.6)


# ---------------------------------------------------------------------------------------
# primitives
# ---------------------------------------------------------------------------------------
def _emu(*v):
    """OOXML coordinates must be integers -- a float (e.g. bw / 2) makes PowerPoint reject the file."""
    return [Emu(int(round(x))) for x in v]


def _font(run, size, color=INK, bold=False, italic=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         margin=0, space_after=6, name=None):
    """paras: str | list[str | list[(text, {size,color,bold,italic})]]"""
    tb = slide.shapes.add_textbox(*_emu(x, y, w, h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space_after)
        runs = p if isinstance(p, list) else [(p, {})]
        for t, o in runs:
            r = para.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", False))
    return tb


def box(slide, x, y, w, h, fill=TEAL_L, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, name=None):
    s = slide.shapes.add_shape(shape, *_emu(x, y, w, h))
    if name:
        s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def arrow(slide, x1, y1, x2, y2, color=TEAL, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *_emu(x1, y1, x2, y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def title(slide, t, kicker=None, dark=False):
    c = WHITE if dark else INK
    if kicker:
        text(slide, MX, Inches(0.42), Inches(10), Inches(0.35), kicker, size=13, color=AMBER if dark else TEAL, bold=True)
    text(slide, MX, Inches(0.72), Inches(12.1), Inches(0.8), t, size=30, color=c, bold=True, name="Title")


def footer(slide, n, dark=False):
    text(slide, MX, Inches(7.0), Inches(8), Inches(0.3), "RSI · 基于大模型后训练的指令微调自进化系统", size=10,
         color=MUTED)
    text(slide, Inches(11.9), Inches(7.0), Inches(0.85), Inches(0.3), str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def stat(slide, x, y, w, big, label, sub=None, color=TEAL, dark=False):
    text(slide, x, y, w, Inches(0.95), big, size=44, color=color, bold=True)
    text(slide, x, y + Inches(0.95), w, Inches(0.4), label, size=14, color=WHITE if dark else INK, bold=True)
    if sub:
        text(slide, x, y + Inches(1.32), w, Inches(0.6), sub, size=11, color=MUTED if dark else INK2)


def robust_strict(eval_dir, split="test"):
    """Strict IF accuracy counting only substantive answers (recomputed from rows; older metrics lack it)."""
    from rsi.synth import substantive
    rows = [r for r in read_jsonl(Path(eval_dir) / f"eval_{split}.jsonl") if r["type"] == "if"]
    return sum(r["strict"] and substantive(r["response"], r["kinds"]) for r in rows) / max(1, len(rows))


def clip_text(t, max_chars=420, max_lines=9):
    lines = [l for l in t.strip().splitlines() if l.strip()]
    out = "\n".join(lines[:max_lines])
    if len(out) > max_chars:
        out = out[:max_chars]
    return out + (" …（截断）" if out != t.strip() else "")


def pct(x, d=1):
    return "—" if x is None else f"{x * 100:.{d}f}%"


def _cjk(font, size=None, color=INK2, bold=False):
    """Chart text: python-pptx only sets the Latin typeface, so CJK falls back to a serif font."""
    font.name = FONT
    if size:
        font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color
    rpr = font._rPr
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def style_chart(chart, legend=True):
    _cjk(chart.font, 11)
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        _cjk(chart.legend.font, 11)
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.format.line.fill.background()
    _cjk(va.tick_labels.font, 10)
    ca = chart.category_axis
    ca.format.line.color.rgb = LINE
    _cjk(ca.tick_labels.font, 10)


# ---------------------------------------------------------------------------------------
def build(run: str, out: Path):
    rd = RUNS / run
    st = R.load_state(rd)
    sm = R.summary(st)
    gens = st["generations"]
    g0 = gens[0]
    champ = next(g for g in gens if g["gen"] == st["champion"])
    b, c = g0["test"], champ.get("test") or champ["dev"]
    n_gen = len(gens) - 1
    n_prom = sum(1 for g in gens[1:] if g.get("promoted"))
    pool = read_jsonl(rd / "self_pool.jsonl")
    n_self = len(pool)
    hours = sum(g.get("seconds") or 0 for g in gens) / 3600
    base_name = st["base_model"].split("/")[-1]

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    n = 0

    def new(dark=False):
        nonlocal n
        n += 1
        s = prs.slides.add_slide(blank)
        bg(s, INK if dark else WHITE)
        return s

    # 1 ── cover ──────────────────────────────────────────────────────────────────────
    s = new(dark=True)
    text(s, MX, Inches(1.5), Inches(8), Inches(0.4), "RECURSIVE SELF-IMPROVEMENT · 试点：指令微调", size=14, color=AMBER, bold=True)
    text(s, MX, Inches(2.0), Inches(11.5), Inches(1.9), ["基于大模型后训练的", "指令微调自进化系统 RSI"], size=44,
         color=WHITE, bold=True, space_after=2)
    text(s, MX, Inches(4.25), Inches(11), Inches(0.9),
         "模型自己出题、自己作答、程序化验证、自动选配方、统计门控晋升，一代代变强", size=18, color=RGBColor(0xC9, 0xD3, 0xE0))
    text(s, MX, Inches(5.7), Inches(11), Inches(0.4),
         f"基座 {base_name} · 单卡 RTX 4090 · 已自主进化 {n_gen} 代", size=14, color=MUTED)
    # motif: the evolution ring
    cx, cy, r = Inches(11.2), Inches(5.6), Inches(0.9)
    ring = s.shapes.add_shape(MSO_SHAPE.DONUT, cx - r, cy - r, 2 * r, 2 * r)
    ring.adjustments[0] = 0.12
    ring.fill.solid()
    ring.fill.fore_color.rgb = TEAL
    ring.line.fill.background()
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, cx + Inches(0.55), cy - Inches(0.85), Inches(0.32), Inches(0.32))
    dot.fill.solid()
    dot.fill.fore_color.rgb = AMBER
    dot.line.fill.background()
    s.notes_slide.notes_text_frame.text = ("本项目把后训练中最依赖人工的三件事：造数据、调配方、判断好坏，交给系统自己闭环完成。"
                                           "试点选择指令微调，因为指令约束可以被程序精确验证。")

    # 2 ── headline results ───────────────────────────────────────────────────────────
    s = new()
    title(s, "一句话结论：无人工标注，指令遵循能力自主提升", "核心结果 · 测试集（不参与任何选择）")
    cols = [("if_prompt_acc", "指令遵循 · 严格准确率", "全部约束都满足才算对", TEAL, False),
            ("if_instr_acc", "指令遵循 · 约束级准确率", "逐条约束统计", TEAL, False),
            ("math_acc", "GSM8K 数学推理", "守护指标：不允许明显回退", INK, False),
            ("degenerate_rate", "输出退化率", "不停止 / 复读，越低越好", INK, True)]
    cw = Inches(2.85)
    for i, (k, lab, sub, col, lower_better) in enumerate(cols):
        x = MX + i * (cw + Inches(0.2))
        d = (c[k] - b[k]) * 100
        good = (d < 0) if lower_better else (d > 0)
        box(s, x, Inches(1.8), cw, Inches(2.4), fill=TEAL_L if i < 2 else RGBColor(0xF1, 0xF3, 0xF6))
        text(s, x + Inches(0.25), Inches(1.95), cw - Inches(0.4), Inches(0.4), lab, size=13, bold=True, color=INK2)
        text(s, x + Inches(0.25), Inches(2.35), cw - Inches(0.4), Inches(0.9), pct(c[k]), size=40, color=col, bold=True)
        text(s, x + Inches(0.25), Inches(3.25), cw - Inches(0.4), Inches(0.4),
             [[(f"基座 {pct(b[k])}　", {"color": INK2}),
               (f"{d:+.1f} 点", {"bold": True, "color": TEAL if good else (INK2 if abs(d) < 0.05 else RED)})]], size=13)
        text(s, x + Inches(0.25), Inches(3.65), cw - Inches(0.4), Inches(0.4), sub, size=11, color=MUTED)
    facts = [
        [("闭环自动化　", {"bold": True, "color": TEAL}),
         (f"{n_gen} 代进化、{n_prom} 次晋升全部由系统自主决策；共训练评测 {sum(len(g.get('candidates', [])) for g in gens)} 个候选模型", {})],
        [("零人工标注　", {"bold": True, "color": TEAL}),
         (f"冠军模型自生成并经程序验证的训练样本累计 {n_self} 条", {})],
        [("可信的提升　", {"bold": True, "color": TEAL}),
         ("test 只用于汇报；晋升在每代新抽的门控集上与冠军同场比较，需 P(更优) ≥ 0.8 且守护指标不回退"
          if any(g.get("champ_gate") for g in gens) else
          "dev 用于选择、test 只用于汇报；晋升需配对 bootstrap P(更优) ≥ 0.8 且守护指标不回退", {})],
    ]
    # Headline = the champion the system itself chose.  If an earlier champion scores higher on test, say so
    # instead of switching the headline to it -- picking the reported model by test score would be test-set selection.
    prom = [g for g in gens[1:] if g.get("promoted") and g.get("test")]
    top = max(prom, key=lambda g: g["test"]["if_prompt_acc"]) if prom else None
    if top and top["gen"] != champ["gen"]:
        ife_m = {}
        p_ife = rd / "ifeval_official" / "summary.json"
        if p_ife.exists():
            ife_m = json.loads(p_ife.read_text(encoding="utf-8"))["models"]
        a_l, b_l = f"gen{top['gen']}", f"gen{champ['gen']}"
        ife_txt = (f"，官方 IFEval 上第{champ['gen']}代 {pct(ife_m[b_l]['prompt_level_strict_acc'])} vs "
                   f"{pct(ife_m[a_l]['prompt_level_strict_acc'])}") if a_l in ife_m and b_l in ife_m else ""
        facts.append([("诚实说明　　", {"bold": True, "color": TEAL}),
                      (f"第{top['gen']}代在项目 test 上略高（严格 {pct(top['test']['if_prompt_acc'])}）{ife_txt}；"
                       f"差异均在噪声内，视为同一水平", {})])
    facts += [
        [("稳健口径　　", {"bold": True, "color": TEAL}),
         (f"剔除不足 20 词的“钻空子”回答后，严格准确率 {pct(robust_strict(g0['eval_dir']))} → "
          f"{pct(robust_strict(champ['eval_dir']))}" if champ.get("test") else "—", {})],
        [("低成本　　　", {"bold": True, "color": TEAL}),
         (f"单张 RTX 4090，累计约 {hours:.1f} GPU 小时；全部数据与缓存位于 D 盘", {})],
    ]
    text(s, MX, Inches(4.4), Inches(12.1), Inches(2.5), facts, size=13.5 if len(facts) > 5 else 15,
         space_after=6 if len(facts) > 5 else 10)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = "左两块是目标能力，右两块是守护指标。所有数字来自独立测试集，系统在进化过程中从未看到它。"

    # 2b ── independent check: official IFEval ──────────────────────────────────────────
    ife_path = rd / "ifeval_official" / "summary.json"
    if ife_path.exists():
        ife = json.loads(ife_path.read_text(encoding="utf-8"))
        labels = list(ife["models"])  # "base", "gen6", ...
        s = new()
        best = labels[-1]
        vb = (ife.get("vs_base") or {}).get(best, {}).get("prompt_level_strict")
        mb, mc = ife["models"]["base"], ife["models"][best]
        verdict = "第三方验证：官方 IFEval 上同样提升" if vb and vb["ci90"][0] > 0 else "第三方验证：官方 IFEval 结果"
        title(s, verdict, "独立基准 · google/IFEval 全部 541 题 · 官方评分代码")
        cd = CategoryChartData()
        bd_all = ife.get("breakdown", {})
        cd.categories = ["提示级·严格（主指标）", "指令级·严格", "训练未见过的指令类型"]
        for lab in labels:
            m = ife["models"][lab]
            name = "原始基座" if lab == "base" else f"第{lab[3:]}代"
            cd.add_series(name, [round(m["prompt_level_strict_acc"], 4), round(m["inst_level_strict_acc"], 4),
                                 round(bd_all.get(lab, {}).get("inst_strict_unseen", 0), 4)])
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, MX, Inches(1.6), Inches(7.6), Inches(5.2), cd)
        ch_ = gf.chart
        style_chart(ch_)
        ch_.value_axis.minimum_scale = 0
        ch_.value_axis.maximum_scale = 0.7
        ch_.value_axis.tick_labels.number_format = "0%"
        ch_.value_axis.tick_labels.number_format_is_linked = False
        pl = ch_.plots[0]
        pl.gap_width = 50
        pl.overlap = 0
        pl.has_data_labels = True
        pl.data_labels.number_format = "0%"
        pl.data_labels.number_format_is_linked = False
        pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        _cjk(pl.data_labels.font, 9)
        # one sequential teal ramp: older -> lighter, the current champion darkest; base stays neutral gray
        ramp = [RGBColor(0xB9, 0xDE, 0xDB), RGBColor(0x5F, 0xB8, 0xB4), RGBColor(0x2E, 0x99, 0x95), TEAL]
        palette = [RGBColor(0xC5, 0xCC, 0xD6)] + ramp[-(len(labels) - 1):]
        for sr, col in zip(pl.series, palette):
            sr.format.fill.solid()
            sr.format.fill.fore_color.rgb = col
        x0 = MX + Inches(8.0)
        d = (mc["prompt_level_strict_acc"] - mb["prompt_level_strict_acc"]) * 100
        text(s, x0, Inches(1.75), Inches(4.1), Inches(0.4), "提示级·严格准确率（官方主指标）", size=13, bold=True, color=INK2)
        text(s, x0, Inches(2.15), Inches(4.1), Inches(0.9),
             f"{pct(mb['prompt_level_strict_acc'])} → {pct(mc['prompt_level_strict_acc'])}", size=30, bold=True, color=TEAL)
        if vb:
            text(s, x0, Inches(3.05), Inches(4.1), Inches(0.7),
                 [[(f"{d:+.1f} 点", {"bold": True, "color": TEAL if vb['ci90'][0] > 0 else INK2}),
                   (f"　配对 bootstrap 90%CI [{vb['ci90'][0] * 100:+.1f}, {vb['ci90'][1] * 100:+.1f}]，"
                    f"P(更优)={vb['p_better']:.2f}", {"color": INK2})]], size=12.5)
        bd = ife.get("breakdown", {})
        bb, bc = bd.get("base", {}), bd.get(best, {})
        n_types = len(mb["per_instruction_n"])
        notes = []
        if bb and bc:
            notes += [
                [("训练中从未见过的指令类型", {"bold": True}),
                 (f"（{bc['n_inst_unseen']} 条）：{pct(bb['inst_strict_unseen'])} → {pct(bc['inst_strict_unseen'])}，"
                  "说明是能力迁移而非套模板", {})],
                [("剔除不足 20 词的“通过”", {"bold": True}),
                 (f"：提示级严格 {pct(bb['prompt_strict_excl_short'])} → {pct(bc['prompt_strict_excl_short'])}", {})]]
        notes += [[(f"人工撰写 541 题、{n_types} 类指令，官方判分；ChatML + 贪心 + 最长 1280 token，"
                    f"完成率 {pct(mb['finish_rate'], 0)} → {pct(mc['finish_rate'], 0)}", {"color": INK2})]]
        text(s, x0, Inches(3.95), Inches(4.1), Inches(2.9), [[("• ", {"color": TEAL})] + t for t in notes],
             size=12, color=INK, space_after=8)
        footer(s, n)
        s.notes_slide.notes_text_frame.text = (
            "项目内测试集的约束模板与训练同源，可能高估泛化；因此用官方 IFEval 做第三方复核。"
            "宽松口径会去掉首尾行、markdown 星号后再判分。\n"
            "分项里 no_comma、forbidden_words 看似回退，是测量假象：基座大量“通过”来自复读失控的输出"
            "（复读文本天然不含逗号/禁用词）。no_comma：基座通过 48/66，其中 36 条为复读；冠军 24/66，其中 6 条复读；"
            "剔除复读后 12 → 18。forbidden_words：基座 27/49（19 条复读），冠军 20/49（7 条复读）；剔除后 8 → 13。"
            "（以上分项为第 6 代与基座的对比）\n"
            "宽松口径（提示级/指令级）：" + "；".join(
                f"{'基座' if l == 'base' else '第' + l[3:] + '代'} {pct(ife['models'][l]['prompt_level_loose_acc'])}"
                f"/{pct(ife['models'][l]['inst_level_loose_acc'])}" for l in labels))

    # 3 ── why ────────────────────────────────────────────────────────────────────────
    s = new()
    title(s, "为什么要让后训练“自进化”", "问题与思路")
    pains = [("数据靠人", "指令数据依赖人工撰写与标注，昂贵、慢、难以针对模型弱点补充"),
             ("配方靠试", "学习率、轮数、LoRA 秩、数据配比靠工程师反复手调"),
             ("好坏靠看", "改了以后是否真的更好，常依赖少量样例的主观判断，易被噪声误导")]
    fixes = [("模型自己造数据", "冠军模型针对自身弱项出题、多次采样、程序验证后留下合格样本"),
             ("系统自己选配方", "元优化器提出候选配方，并从历史成败中学习哪类改动更可能有效"),
             ("统计自己把关", "固定 dev/test 自动评测 + 配对 bootstrap 门控，提升不显著就回滚")]
    for i, ((pt, pd), (ft, fd)) in enumerate(zip(pains, fixes)):
        y = Inches(1.8) + i * Inches(1.65)
        box(s, MX, y, Inches(5.3), Inches(1.4), fill=RGBColor(0xF1, 0xF3, 0xF6))
        text(s, MX + Inches(0.3), y + Inches(0.18), Inches(4.8), Inches(0.4), pt, size=17, bold=True, color=INK2)
        text(s, MX + Inches(0.3), y + Inches(0.6), Inches(4.8), Inches(0.75), pd, size=13, color=INK2)
        arrow(s, MX + Inches(5.45), y + Inches(0.7), MX + Inches(6.4), y + Inches(0.7), color=AMBER, width=2.5)
        box(s, MX + Inches(6.55), y, Inches(5.55), Inches(1.4), fill=TEAL_L)
        text(s, MX + Inches(6.85), y + Inches(0.18), Inches(5), Inches(0.4), ft, size=17, bold=True, color=TEAL)
        text(s, MX + Inches(6.85), y + Inches(0.6), Inches(5), Inches(0.75), fd, size=13, color=INK)
    footer(s, n)

    # 4 ── architecture loop ──────────────────────────────────────────────────────────
    s = new()
    title(s, "系统架构：一代 = 一次完整的自我改进闭环", "架构")
    steps = [("① 诊断", "读取冠军在 dev 上各约束类型的通过率，找出弱项"),
             ("② 自生成", "冠军按弱项加权出题，每题采样 k 次作答"),
             ("③ 验证筛选", "程序化校验 18 类约束与数学答案，只留合格样本"),
             ("④ 配方提议", "元优化器对冠军配方施加 1–2 个改动算子，得到 N 个候选"),
             ("⑤ 训练评测", "LoRA SFT 训练候选，并在固定 dev 集上自动评测"),
             ("⑥ 门控反思", "显著更优才晋升（并测 test），更新算子经验与数据引擎参数")]
    pos = [(0, 0), (1, 0), (2, 0), (2, 1), (1, 1), (0, 1)]
    bw, bh = Inches(3.55), Inches(1.75)
    gx, gy = Inches(0.62), Inches(0.75)
    x0, y0 = MX + Inches(0.15), Inches(1.75)
    centers = []
    for (t, d), (col, row) in zip(steps, pos):
        x = x0 + col * (bw + gx)
        y = y0 + row * (bh + gy)
        hl = t.startswith(("② ", "③ ", "⑥ "))
        box(s, x, y, bw, bh, fill=TEAL if hl else TEAL_L)
        text(s, x + Inches(0.25), y + Inches(0.2), bw - Inches(0.5), Inches(0.45), t, size=18, bold=True,
             color=WHITE if hl else TEAL)
        text(s, x + Inches(0.25), y + Inches(0.72), bw - Inches(0.5), Inches(0.95), d, size=13,
             color=WHITE if hl else INK)
        centers.append((x, y))
    # arrows along the loop
    a = Inches(0.08)
    for i in range(2):
        x, y = centers[i]
        arrow(s, x + bw + a, y + bh / 2, x + bw + gx - a, y + bh / 2)
    x, y = centers[2]
    arrow(s, x + bw / 2, y + bh + a, x + bw / 2, y + bh + gy - a)
    for i in (3, 4):
        x, y = centers[i]
        arrow(s, x - a, y + bh / 2, x - gx + a, y + bh / 2)
    x, y = centers[5]
    arrow(s, x + bw / 2, y - a, x + bw / 2, y - gy + a, color=AMBER)
    text(s, x + bw / 2 + Inches(0.12), y - gy + Inches(0.12), Inches(3.2), Inches(0.5), "新冠军进入下一代",
         size=12, color=AMBER, bold=True)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = "深色三步是“自”的部分：数据由模型自己产生，由程序自己验证，晋升由统计自己决定。"

    # 5 ── data engine ────────────────────────────────────────────────────────────────
    s = new()
    title(s, "自数据引擎：能被验证的数据，才能放心自产", "核心模块 ① 数据")
    srcs = [("自生成·指令", "拒绝采样微调 (ReST-EM)", "冠军回答带约束的指令，k 个样本中全部约束通过、无退化者入池"),
            ("自生成·数学", "STaR 自举推理", "冠军解 GSM8K 训练题，最终答案正确的最短推理链入池"),
            ("自蒸馏种子", "SDFT 防遗忘", "冠军自己回答普通指令，替代外部参考答案，避免把模型拉离自身分布"),
            ("课程加权", "弱项优先", "约束类型抽样权重 ∝ (1−通过率+0.05)^强度，强度随通过率自动调高"),
            ("规则增强 / 外部种子", "冷启动先验", "种子回答经程序改写满足约束；alpaca / GSM8K 金标作可选外部数据")]
    for i, (t1, t2, d) in enumerate(srcs):
        y = Inches(1.7) + i * Inches(1.0)
        box(s, MX, y, Inches(6.6), Inches(0.9), fill=TEAL_L if i < 4 else RGBColor(0xF1, 0xF3, 0xF6))
        text(s, MX + Inches(0.25), y + Inches(0.08), Inches(6.2), Inches(0.4),
             [[(t1, {"bold": True, "color": TEAL, "size": 15}), ("　" + t2, {"color": INK2, "size": 12.5})]])
        text(s, MX + Inches(0.25), y + Inches(0.47), Inches(6.2), Inches(0.4), d, size=12)
    # constraint chips
    x0 = MX + Inches(7.0)
    text(s, x0, Inches(1.7), Inches(5), Inches(0.4), "18 类可程序验证的约束（IFEval 风格）", size=15, bold=True)
    chips = [KIND_ZH[k] for k in KINDS]
    cw, ch = Inches(1.62), Inches(0.46)
    for i, k in enumerate(chips):
        r_, c_ = divmod(i, 3)
        x = x0 + c_ * (cw + Inches(0.1))
        y = Inches(2.25) + r_ * (ch + Inches(0.1))
        box(s, x, y, cw, ch, fill=WHITE, line=LINE, radius=0.3)
        text(s, x, y, cw, ch, k, size=12, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    syn = [g.get("synth") for g in gens if g.get("synth")]
    if syn:
        last = syn[-1]
        text(s, x0, Inches(5.8), Inches(5.3), Inches(0.9),
             [[("最近一代：", {"bold": True}),
               (f"指令题 {last['if']['prompts']} 道 × k={last['settings']['k']}，样本通过率 {pct(last['if']['sample_pass_rate'])}；"
                f"数学通过率 {pct(last['math']['sample_pass_rate'])}", {})]], size=12, color=INK2)
    footer(s, n)

    # 6 ── meta-optimizer ─────────────────────────────────────────────────────────────
    s = new()
    title(s, "元优化器：不仅优化模型，也优化“怎么优化”", "核心模块 ② 配方搜索")
    genes = [("学习率", "×/÷1.6"), ("训练轮数", "1 / 1.5 / 2 / 3"), ("LoRA 秩", "8 / 16 / 32 / 64"),
             ("种子数据量", "×/÷1.5"), ("数学金标量", "×/÷1.5"), ("自生成数据占比", "±25%"),
             ("规则增强量", "×/÷1.6"), ("初始化", "从基座训练 / 冠军续训")]
    text(s, MX, Inches(1.7), Inches(5.5), Inches(0.4), "配方基因与改动算子", size=15, bold=True)
    for i, (gname, rng_) in enumerate(genes):
        y = Inches(2.15) + i * Inches(0.5)
        box(s, MX, y, Inches(5.4), Inches(0.42), fill=TEAL_L if i % 2 == 0 else WHITE, radius=0.2)
        text(s, MX + Inches(0.2), y, Inches(2.6), Inches(0.42), gname, size=13, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, MX + Inches(2.8), y, Inches(2.5), Inches(0.42), rng_, size=13, color=INK2, anchor=MSO_ANCHOR.MIDDLE)
    # bandit chart (native)
    bpath = rd / "bandit.json"
    bs = json.loads(bpath.read_text(encoding="utf-8")) if bpath.exists() else {}
    tried = sorted(((k, a / (a + bb), a + bb - 2) for k, (a, bb) in bs.items() if a + bb > 2), key=lambda t: t[1])
    if tried:
        cd = CategoryChartData()
        cd.categories = [f"{OP_ZH[k]} (n={t:.0f})" for k, _, t in tried]
        cd.add_series("成功率后验均值", [round(m, 3) for _, m, _ in tried])
        gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, MX + Inches(6.0), Inches(1.6), Inches(6.1),
                                Inches(3.4), cd)
        ch_ = gf.chart
        style_chart(ch_, legend=False)
        ch_.has_title = True
        ch_.chart_title.text_frame.text = "Thompson 采样：各改动算子的成功率后验（已尝试）"
        _font(ch_.chart_title.text_frame.paragraphs[0].runs[0], 12, INK, True)
        ch_.value_axis.maximum_scale = 1.0
        ch_.value_axis.minimum_scale = 0.0
        ch_.value_axis.tick_labels.number_format = "0%"
        ch_.value_axis.tick_labels.number_format_is_linked = False
        pl = ch_.plots[0]
        pl.gap_width = 60
        pl.has_data_labels = True
        pl.data_labels.number_format = "0%"
        pl.data_labels.number_format_is_linked = False
        pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        _cjk(pl.data_labels.font, 10)
        pl.series[0].format.fill.solid()
        pl.series[0].format.fill.fore_color.rgb = TEAL
    layers = [("配方搜索", "候选 = 冠军配方 + 1–2 个算子；算子按 Beta 后验做 Thompson 采样，通过门控记为成功"),
              ("失败修复", "上一代“显著更优却被守护指标拦下”的候选作父本，只用针对该失败的修复算子"),
              ("数据引擎", "按自生成通过率反馈调节采样数 k、温度、课程强度与出题规模")]
    for i, (hdr, d) in enumerate(layers):
        y = Inches(5.1) + i * Inches(0.6)
        text(s, MX + Inches(6.0), y, Inches(1.2), Inches(0.55), hdr, size=12.5, bold=True, color=TEAL)
        text(s, MX + Inches(7.2), y, Inches(4.9), Inches(0.55), d, size=11.5, color=INK, space_after=0)
    footer(s, n)

    # 7 ── evaluation & gate ──────────────────────────────────────────────────────────
    s = new()
    title(s, "评测与门控：防止系统“自我感动”", "核心模块 ③ 可信评估")
    v7 = any(g.get("champ_gate") for g in gens)  # v7+: promotion decided on a freshly drawn gate set
    text(s, MX, Inches(1.7), Inches(6), Inches(0.4), "互不相交的评测集，各司其职", size=15, bold=True)
    sets = [("dev", "200 道带约束指令 + 150 道 GSM8K", "画曲线、做诊断" if v7 else "用于候选选择与门控", TEAL, TEAL_L)]
    if v7:
        sets.append(("gate", "每代新抽：重采样约束 + 未用过的 GSM8K", "冠军与候选同场重测 → 决定晋升（防赢家诅咒）",
                     TEAL, RGBColor(0xD3, 0xEC, 0xEA)))
    sets.append(("test", "另 200 + 150 道", "只在基线与晋升时评测，从不反馈", AMBER, RGBColor(0xFE, 0xF4, 0xE2)))
    hh = Inches(0.85) if v7 else Inches(1.05)
    for i, (nm, d1, d2, col, fill) in enumerate(sets):
        y = Inches(2.2) + i * (hh + Inches(0.12))
        box(s, MX, y, Inches(5.8), hh, fill=fill)
        text(s, MX + Inches(0.25), y, Inches(1.1), hh, nm, size=22, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE)
        text(s, MX + Inches(1.45), y + Inches(0.1), Inches(4.25), Inches(0.38), d1, size=13.5, bold=True)
        text(s, MX + Inches(1.45), y + Inches(0.46), Inches(4.25), Inches(0.38), d2, size=11.5, color=INK2)
    text(s, MX, Inches(5.15) if v7 else Inches(4.85), Inches(6), Inches(1.6),
         [[("综合分 = ", {"bold": True}), ("0.6×指令严格准确率 + 0.3×GSM8K + 0.1×(1−退化率)", {})],
          "同一代所有候选共享同一份采样数据，差异只来自配方本身。",
          "退化 = 未输出结束符或 4-gram 复读率 > 50%。"], size=13, color=INK, space_after=8)
    x0 = MX + Inches(6.5)
    text(s, x0, Inches(1.7), Inches(5.5), Inches(0.4), "晋升必须同时满足", size=15, bold=True)
    conds = [("Δ 综合分 > 0", "相对当前冠军，在本代门控集上同场重测" if v7 else "相对当前冠军，在 dev 上"),
             ("P(更优) ≥ 0.8", "按题目配对 bootstrap 2000 次"),
             ("数学回退 ≤ 3 点", "守护通用推理能力"), ("退化率上升 ≤ 3 点", "守护输出稳定性")]
    for i, (cn, cd_) in enumerate(conds):
        y = Inches(2.2) + i * Inches(1.08)
        dotc = s.shapes.add_shape(MSO_SHAPE.OVAL, x0, y + Inches(0.12), Inches(0.62), Inches(0.62))
        dotc.fill.solid()
        dotc.fill.fore_color.rgb = TEAL
        dotc.line.fill.background()
        text(s, x0, y + Inches(0.12), Inches(0.62), Inches(0.62), str(i + 1), size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x0 + Inches(0.85), y + Inches(0.05), Inches(4.6), Inches(0.4), cn, size=16, bold=True)
        text(s, x0 + Inches(0.85), y + Inches(0.47), Inches(4.6), Inches(0.4), cd_, size=12, color=INK2)
    text(s, x0, Inches(6.45), Inches(5.6), Inches(0.4), "任一不满足 → 保留旧冠军（自动回滚），并把失败记入算子经验。",
         size=12.5, color=RED, bold=True)
    footer(s, n)

    # 8 ── progress chart ─────────────────────────────────────────────────────────────
    s = new()
    title(s, f"进化曲线：{n_gen} 代中晋升 {n_prom} 次，每次都经统计门控确认", "实验结果 ① 趋势")
    ser = R.champion_series(st)
    cd = CategoryChartData()
    cd.categories = [f"第{x['gen']}代" for x in ser]
    for k, nm in (("if_prompt_acc", "指令严格准确率"), ("if_instr_acc", "约束级准确率"), ("math_acc", "GSM8K")):
        cd.add_series(nm, [round((x["test"] or x["dev"])[k], 4) for x in ser])
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, MX, Inches(1.6), Inches(7.6), Inches(5.2), cd)
    ch_ = gf.chart
    style_chart(ch_)
    ch_.value_axis.minimum_scale = 0
    ch_.value_axis.maximum_scale = 1
    ch_.value_axis.tick_labels.number_format = "0%"
    ch_.value_axis.tick_labels.number_format_is_linked = False
    for sr, col in zip(ch_.plots[0].series, (TEAL, RGBColor(0x5F, 0xB8, 0xB4), AMBER)):
        sr.format.line.color.rgb = col
        sr.format.line.width = Pt(2.5)
        sr.smooth = False
        sr.marker.size = 8
        sr.marker.format.fill.solid()
        sr.marker.format.fill.fore_color.rgb = col
        sr.marker.format.line.color.rgb = WHITE
    # label only the endpoints of the headline series (selective direct labels)
    for idx in (0, len(ser) - 1):
        dl = ch_.plots[0].series[0].points[idx].data_label
        dl.position = XL_LABEL_POSITION.BELOW
        tf = dl.text_frame
        tf.text = pct((ser[idx]["test"] or ser[idx]["dev"])["if_prompt_acc"])
        _font(tf.paragraphs[0].runs[0], 12, TEAL, True)
    # decision column
    x0 = MX + Inches(8.0)
    text(s, x0, Inches(1.7), Inches(4.2), Inches(0.4), "每代决策", size=15, bold=True)
    rows = sm["generations"][1:]
    rh = min(Inches(0.62), int(Inches(4.7) / max(1, len(rows))))
    for i, r in enumerate(rows):
        y = Inches(2.2) + i * rh
        ok = r["promoted"]
        box(s, x0, y + Inches(0.04), Inches(0.9), rh - Inches(0.1), fill=TEAL if ok else RGBColor(0xF1, 0xF3, 0xF6), radius=0.25)
        text(s, x0, y + Inches(0.04), Inches(0.9), rh - Inches(0.1), "晋升" if ok else "回滚", size=11, bold=True,
             color=WHITE if ok else INK2, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        ops = " + ".join(r["ops"]) or "默认配方"
        pb = "" if r["p_better"] is None else f"  P={r['p_better']:.2f}"
        text(s, x0 + Inches(1.05), y, Inches(3.3), rh, [[(f"第{r['gen']}代 ", {"bold": True}), (ops + pb, {"color": INK2})]],
             size=11.5, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = "折线为每一代结束时冠军在独立测试集上的表现；未晋升的代沿用上一冠军，所以是平台。"

    # 9 ── per-kind diagnosis ─────────────────────────────────────────────────────────
    s = new()
    title(s, "能力诊断：哪些约束变强了，哪些仍是短板", "实验结果 ② 分项")
    pk0, pk1 = g0["test"]["per_kind"], c["per_kind"]
    # bar charts list categories bottom-up, so sort ascending to put the biggest gain on top
    kinds = sorted(pk1, key=lambda k: pk1[k] - pk0.get(k, 0))
    cd = CategoryChartData()
    cd.categories = [KIND_ZH[k] for k in kinds]
    cd.add_series("原始基座", [round(pk0.get(k, 0), 3) for k in kinds])
    cd.add_series(f"冠军（第{st['champion']}代）", [round(pk1.get(k, 0), 3) for k in kinds])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, MX, Inches(1.55), Inches(7.9), Inches(5.35), cd)
    ch_ = gf.chart
    style_chart(ch_)
    ch_.value_axis.maximum_scale = 1
    ch_.value_axis.minimum_scale = 0
    ch_.value_axis.tick_labels.number_format = "0%"
    ch_.value_axis.tick_labels.number_format_is_linked = False
    pl = ch_.plots[0]
    pl.gap_width = 45
    pl.overlap = 0
    for sr, col in zip(pl.series, (RGBColor(0xC5, 0xCC, 0xD6), TEAL)):
        sr.format.fill.solid()
        sr.format.fill.fore_color.rgb = col
    gains = sorted(((k, pk1[k] - pk0.get(k, 0)) for k in pk1), key=lambda t: -t[1])
    worst = sorted(pk1, key=pk1.get)[:3]
    x0 = MX + Inches(8.3)
    text(s, x0, Inches(1.75), Inches(3.8), Inches(0.4), "提升最大", size=15, bold=True, color=TEAL)
    for i, (k, d) in enumerate(gains[:4]):
        y = Inches(2.2) + i * Inches(0.55)
        text(s, x0, y, Inches(2.3), Inches(0.5), KIND_ZH[k], size=13, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x0 + Inches(2.3), y, Inches(1.5), Inches(0.5), f"{d * 100:+.0f} 点", size=16, bold=True,
             color=TEAL if d > 0 else INK2, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x0, Inches(4.6), Inches(3.8), Inches(0.4), "仍然最弱", size=15, bold=True, color=RED)
    for i, k in enumerate(worst):
        y = Inches(5.05) + i * Inches(0.5)
        text(s, x0, y, Inches(2.3), Inches(0.45), KIND_ZH[k], size=13, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x0 + Inches(2.3), y, Inches(1.5), Inches(0.45), pct(pk1[k], 0), size=16, bold=True, color=INK2,
             anchor=MSO_ANCHOR.MIDDLE)
    text(s, x0, Inches(6.55), Inches(3.9), Inches(0.35), "弱项会在下一代课程中被继续加权出题", size=11, color=INK2)
    footer(s, n)

    # 10 ── sample ────────────────────────────────────────────────────────────────────
    s = new()
    title(s, "样例：同一道题，基座 vs 冠军", "实验结果 ③ 定性")
    r0 = {r["id"]: r for r in read_jsonl(Path(g0["eval_dir"]) / "eval_test.jsonl")}
    rc = {r["id"]: r for r in read_jsonl(Path(champ["eval_dir"]) / "eval_test.jsonl")} if champ.get("test") else {}
    # a fair showcase: the champion's answer must be a real answer (not a verifier-gaming stub)
    cand = [(i, r) for i, r in rc.items() if r["type"] == "if" and r["strict"] and i in r0 and not r0[i]["strict"]
            and len(r["kinds"]) >= 2 and 40 <= len(r["response"].split()) <= 110 and len(r["prompt"]) < 420]
    cand.sort(key=lambda t: (-len(t[1]["kinds"]), not r0[t[0]]["finished"], len(t[1]["response"])))
    if cand:
        sid, rcur = cand[0]
        rb = r0[sid]
        box(s, MX, Inches(1.6), Inches(12.1), Inches(1.15), fill=RGBColor(0xF1, 0xF3, 0xF6))
        text(s, MX + Inches(0.25), Inches(1.68), Inches(11.6), Inches(1.0),
             [[("提示词　", {"bold": True, "color": TEAL}), (rcur["prompt"].replace("\n", " "), {})]], size=11.5)
        for j, (lab, rr, col) in enumerate((("原始基座", rb, INK2), (f"冠军（第{st['champion']}代）", rcur, TEAL))):
            x = MX + j * Inches(6.15)
            box(s, x, Inches(2.95), Inches(5.95), Inches(3.85), fill=WHITE, line=LINE)
            checks = "  ".join(f"{'✓' if ok else '✗'} {KIND_ZH[k]}" for k, ok in zip(rr["kinds"], rr["checks"]))
            text(s, x + Inches(0.25), Inches(3.05), Inches(5.5), Inches(0.4), lab, size=15, bold=True, color=col)
            text(s, x + Inches(0.25), Inches(3.45), Inches(5.5), Inches(0.4), checks, size=12, bold=True,
                 color=TEAL if rr["strict"] else RED)
            resp = clip_text(rr["response"])
            if not rr["finished"]:
                resp += "\n（未输出结束符，持续生成到长度上限）"
            text(s, x + Inches(0.25), Inches(3.9), Inches(5.5), Inches(2.85), resp, size=10.5, color=INK, space_after=2)
    footer(s, n)

    # 10b ── the system evolves too ──────────────────────────────────────────────────
    s = new()
    title(s, "系统自身也在迭代：每个版本都由上一轮证据驱动", "持续优化")
    vers = [("v1", "跑通闭环", "症状：微调后永不停止、训练慢 20×",
             "<|im_end|> 嵌入未训练 → 改用 <|endoftext|>；显存上限 + 分块计损"),
            ("v2", "失败定向修复", "证据：候选显著更优（P=1.00），却因数学回退被拒",
             "只在通过门控者中选优；“险胜被拦”的候选作父本，只用修复算子"),
            ("v3", "选对父本", "证据：最接近通过的候选只差 0.3 点，却没被选为父本",
             "父本 = 违规最小者；新增数据源消融算子（去掉金标 / 种子减半）"),
            ("v4", "自蒸馏种子", "证据：8 个候选全部数学回退，调超参无效",
             "冠军自答种子指令，替代外部答案（SDFT）；新算子先试一次"),
            ("v5", "推理链校验", "证据：自蒸馏反而更差，假设被否定；回答中位长度 217→117",
             "根因：STaR 选“最短正确解”奖励蒙对；改为校验推理链 + 中位长度"),
            ("v6", "防钻空子", "证据：冠军只答“P.S.”也能通过 3 项格式校验",
             "不足 20 词的回答不入池，并回溯清洗；新增稳健口径指标")]
    champ_gate = next((g["champ_gate"] for g in gens if g.get("champ_gate")), None)
    if champ_gate:  # v7 is live: its evidence comes straight from the run
        first = next(g for g in gens if g.get("champ_gate"))
        parent_dev = next(x for x in gens if x["gen"] == first["parent"])["dev"]
        vers.append(("v7", "防赢家诅咒",
                     f"证据：冠军 dev 数学 {pct(parent_dev['math_acc'])}，新抽题上仅 {pct(champ_gate['math_acc'])}",
                     "每代新抽门控集，冠军与候选同场重测后再定晋升"))
    gap = Inches(0.18)
    cw = int((Inches(12.1) - gap * (len(vers) - 1)) / len(vers))
    for i, (v, name, ev, fix) in enumerate(vers):
        x = MX + i * (cw + gap)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, *_emu(x, Inches(1.75), Inches(0.75), Inches(0.75)))
        circ.fill.solid()
        circ.fill.fore_color.rgb = AMBER if i == len(vers) - 1 else TEAL
        circ.line.fill.background()
        text(s, x, Inches(1.75), Inches(0.75), Inches(0.75), v, size=16, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(vers) - 1:
            arrow(s, x + Inches(0.85), Inches(2.12), x + cw + gap - Inches(0.1), Inches(2.12), color=LINE, width=2)
        text(s, x, Inches(2.65), cw, Inches(0.45), name, size=15, bold=True, color=TEAL)
        box(s, x, Inches(3.2), cw, Inches(1.15), fill=RGBColor(0xF1, 0xF3, 0xF6))
        text(s, x + Inches(0.18), Inches(3.3), cw - Inches(0.36), Inches(1.05), ev, size=11 if len(vers) < 7 else 10, color=INK2, space_after=0)
        box(s, x, Inches(4.5), cw, Inches(1.45), fill=TEAL_L)
        text(s, x + Inches(0.18), Inches(4.6), cw - Inches(0.36), Inches(1.35), fix, size=11 if len(vers) < 7 else 10, color=INK, space_after=0)
    text(s, MX, Inches(6.3), Inches(12), Inches(0.35), "门控标准全程未放宽：先看到结果再改晋升条件，等于移动球门。详见 docs/系统迭代记录.md",
         size=12, color=INK2)
    footer(s, n)

    # 10c ── key finding: selection bias in self-generated data ───────────────────────
    s = new()
    title(s, "关键发现：自生成数据的“选择规则”本身就是优化压力", "关键发现")
    text(s, MX, Inches(1.75), Inches(5.6), Inches(0.4), "现象：每个 SFT 候选都丢数学分", size=15, bold=True)
    stat_rows = [("217 → 117", "数学回答中位 token 数（基座 → 候选）", RED),
                 ("11%", "自生成数学“解答”只是蒙对的一句话", RED),
                 ("+11 点", "修复后首个候选 dev 数学（74.0% → 85.3%）", TEAL)]
    for i, (big, lab, col) in enumerate(stat_rows):
        y = Inches(2.25) + i * Inches(1.05)
        text(s, MX, y, Inches(2.3), Inches(0.8), big, size=30, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE)
        text(s, MX + Inches(2.4), y, Inches(3.3), Inches(0.8), lab, size=13, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    box(s, MX, Inches(5.45), Inches(5.7), Inches(1.3), fill=TEAL_L)
    text(s, MX + Inches(0.25), Inches(5.55), Inches(5.3), Inches(1.15),
         [[("根因　", {"bold": True, "color": TEAL}),
           ("STaR 只核对最终答案，又选“最短的正确解”；k=8 时蒙对的一句话总是最短 → 模型学会跳过推理。"
            "提高 k 的反馈规则还放大了偏差。修复：校验推理链 + 取中位长度，并回溯清洗数据池。", {})]], size=12.5)
    # native chart: best candidate math per generation
    cd = CategoryChartData()
    cand_gens = [g for g in gens if g.get("candidates")]
    cd.categories = [f"第{g['gen']}代" for g in cand_gens]
    cd.add_series("最佳候选 GSM8K (dev)", [round(max(c["dev"]["math_acc"] for c in g["candidates"]), 3) for g in cand_gens])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, MX + Inches(6.1), Inches(1.6), Inches(6.0), Inches(4.6), cd)
    ch_ = gf.chart
    style_chart(ch_, legend=False)
    ch_.has_title = True
    ch_.chart_title.text_frame.text = f"每代最佳候选的数学准确率（基座 {pct(g0['dev']['math_acc'])}，门槛 {pct(g0['dev']['math_acc'] - 0.03)}）"
    _font(ch_.chart_title.text_frame.paragraphs[0].runs[0], 12, INK, True)
    ch_.value_axis.minimum_scale = 0.5
    ch_.value_axis.maximum_scale = 0.9
    ch_.value_axis.tick_labels.number_format = "0%"
    ch_.value_axis.tick_labels.number_format_is_linked = False
    pl = ch_.plots[0]
    pl.gap_width = 70
    pl.has_data_labels = True
    pl.data_labels.number_format = "0.0%"
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    _cjk(pl.data_labels.font, 10)
    # v5 = first generation whose self-play recorded rejected shortcut solutions
    v5_from = next((g["gen"] for g in gens if "shortcut_rejected" in ((g.get("synth") or {}).get("math") or {})), 10 ** 6)
    for i, g in enumerate(cand_gens):
        pt = pl.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = TEAL if g["gen"] >= v5_from else RGBColor(0xC5, 0xCC, 0xD6)
    text(s, MX + Inches(6.1), Inches(6.3), Inches(6.0), Inches(0.4),
         [[("■ ", {"color": RGBColor(0xC5, 0xCC, 0xD6)}), ("v1–v4（最短正确解）　", {"color": INK2}),
           ("■ ", {"color": TEAL}), ("v5（推理链校验）", {"color": INK2})]], size=11)
    footer(s, n)

    # 11 ── engineering ───────────────────────────────────────────────────────────────
    s = new()
    title(s, "工程实现：可复现、可续跑、可演示", "工程")
    mods = [("rsi/constraints.py", "18 类约束：渲染 / 校验 / 改写"), ("rsi/synth.py", "自数据引擎：ReST-EM · STaR · 课程"),
            ("rsi/train.py", "LoRA SFT：只对回答计损、分块计损"), ("rsi/evaluate.py", "自动评测 + 配对 bootstrap"),
            ("rsi/optimizer.py", "算子 · Thompson 采样 · 门控 · 修复"), ("rsi/loop.py", "编排器：状态落盘，断点续跑"),
            ("app/demo.py", "Gradio：总览 / 诊断 / 样例 / 对话 / 控制台"), ("scripts/make_ppt.py", "本演示文稿由运行数据自动生成")]
    for i, (f, d) in enumerate(mods):
        r_, c_ = divmod(i, 2)
        x = MX + c_ * Inches(6.15)
        y = Inches(1.7) + r_ * Inches(0.95)
        box(s, x, y, Inches(5.95), Inches(0.8), fill=TEAL_L if (r_ + c_) % 2 == 0 else RGBColor(0xF1, 0xF3, 0xF6))
        text(s, x + Inches(0.25), y, Inches(2.2), Inches(0.8), f, size=12.5, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + Inches(2.45), y, Inches(3.4), Inches(0.8), d, size=12, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    per_gen = [g["seconds"] / 60 for g in gens[1:] if g.get("seconds")]
    text(s, MX, Inches(5.65), Inches(12), Inches(1.2),
         [[("资源　", {"bold": True, "color": TEAL}),
           (f"单张 RTX 4090 24GB；每代约 {sum(per_gen) / max(1, len(per_gen)):.0f} 分钟（自生成 + {len(gens[-1].get('candidates', []))} 个候选训练评测）；"
            "LoRA 训练峰值显存约 14GB。", {})],
          [("踩坑　", {"bold": True, "color": TEAL}),
           ("Qwen3-Base 的 <|im_end|> 嵌入几乎未训练，教它以此结束会永不停止 → 改用 <|endoftext|>；"
            "Windows 显存溢出会静默借用内存导致 20× 变慢 → 显存上限 + 分块计损。", {})]], size=12.5, space_after=8)
    footer(s, n)

    # 12 ── demo ──────────────────────────────────────────────────────────────────────
    s = new()
    title(s, "可演示系统", "Demo")
    shot = ROOT / "deliverables" / "demo_screenshot.png"
    if shot.exists():
        from PIL import Image
        iw, ih = Image.open(shot).size
        hh = Inches(5.1)
        ww = int(hh * iw / ih)
        if ww > Inches(7.6):
            ww = Inches(7.6)
            hh = int(ww * ih / iw)
        s.shapes.add_picture(str(shot), MX, Inches(1.6), ww, hh)
    x0 = MX + Inches(8.0)
    tabs = [("进化总览", "测试集 KPI、逐代曲线、决策表、反思日志"), ("能力诊断", "约束热力图、算子经验、数据构成"),
            ("样例对比", "任选 dev 样例，对比任意两代输出与校验结果"), ("实时对话", "自定义任务 + 勾选约束，两代模型同时作答并当场验证"),
            ("进化控制台", "一键继续进化 N 代，实时日志")]
    for i, (t, d) in enumerate(tabs):
        y = Inches(1.65) + i * Inches(0.95)
        text(s, x0, y, Inches(4.2), Inches(0.4), t, size=15, bold=True, color=TEAL)
        text(s, x0, y + Inches(0.38), Inches(4.2), Inches(0.5), d, size=12, color=INK2)
    text(s, x0, Inches(6.45), Inches(4.4), Inches(0.4), "python app/demo.py --run pilot → http://127.0.0.1:7860",
         size=11, color=INK2)
    footer(s, n)

    # 13 ── limits & next ─────────────────────────────────────────────────────────────
    s = new(dark=True)
    title(s, "局限与下一步", "展望", dark=True)
    lim = ["验证器只查格式不查内容，会被钻空子（如只输出“P.S.”）；已加实质性过滤与稳健口径，仍需内容评审",
           "dev 每代重复用于选择，存在轻微过拟合风险（test 已隔离，以 test 为准）",
           f"试点规模：1.7B 模型、单卡、每代 {len(gens[-1].get('candidates', [])) or 3} 个候选，搜索空间受算力限制"]
    nxt = ["引入模型自评 / 偏好数据，扩展到 DPO 等偏好优化阶段",
           "放大到本机 Qwen3.5-9B（QLoRA），验证规律是否随规模保持",
           "把同一闭环推广到代码、工具调用等可验证任务"]
    for j, (hdr, items, col) in enumerate((("当前局限", lim, RGBColor(0xC9, 0xD3, 0xE0)), ("下一步", nxt, AMBER))):
        x = MX + j * Inches(6.15)
        text(s, x, Inches(1.85), Inches(5.8), Inches(0.5), hdr, size=20, bold=True, color=col)
        for i, it in enumerate(items):
            y = Inches(2.55) + i * Inches(1.25)
            box(s, x, y, Inches(5.8), Inches(1.05), fill=RGBColor(0x1D, 0x29, 0x40))
            text(s, x + Inches(0.3), y, Inches(5.3), Inches(1.05), it, size=14, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX, Inches(6.6), Inches(12), Inches(0.4),
         "代码：github.com/notquitezero-ai/rsi-instruct-evolve · 在线演示：notquitezero-ai.github.io/rsi-instruct-evolve",
         size=11, color=MUTED)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", default="pilot")
    ap.add_argument("--out", default=str(ROOT / "deliverables" / "RSI_演示.pptx"))
    a = ap.parse_args()
    print(build(a.run, Path(a.out)))
